"""Lecture des rapports Excel de Sécurité Incendie National (migration de
leur historique — voir la commande `importer_securite_nationale`).

Quatre gabarits, un bâtiment et une visite par classeur :
- « ALARME S-536 » : fiches E1, E2.x, légende et E3.2.1 de la norme ;
- « EXTINCTEURS NFPA 10 » : tableau des extincteurs portatifs ;
- « ÉCLAIRAGE NFPA 101 » : tableau des unités d'éclairage d'urgence ;
- « GICLEURS NFPA 25 » : formulaire d'une page (vannes, tuyauterie,
  raccord pompier, système sous eau, pompes, alimentation d'eau).

Ce module ne touche pas à la base : il transforme un classeur en dict
neutre et note les anomalies (`avertissements`) au lieu de deviner.
Les cellules sont repérées par leur libellé plutôt que par leur adresse,
pour tolérer les petits décalages d'une copie à l'autre du gabarit.
"""

import re
import unicodedata
from datetime import date, datetime, time

from openpyxl import load_workbook

ALARME = "alarme"
EXTINCTEURS = "extincteurs"
ECLAIRAGE = "eclairage"
GICLEURS = "gicleurs"

MOIS = {
    "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6, "juillet": 7,
    "aout": 8, "septembre": 9, "octobre": 10, "novembre": 11, "decembre": 12,
}

# Valeurs laissées dans le gabarit quand l'information est inconnue.
VIDES = {"", "xx", "x x", "0", "-", "—", "n/a", "na"}


# ── Utilitaires ──────────────────────────────────────────────────────────────

def norm(texte) -> str:
    """Minuscules, sans accents, espaces fusionnés, sans « : » final."""
    if texte is None:
        return ""
    texte = unicodedata.normalize("NFKD", str(texte).replace("’", "'")).encode("ascii", "ignore").decode()
    texte = re.sub(r"\s+", " ", texte).strip().lower()
    return texte.rstrip(" :").strip()


def propre(valeur) -> str:
    """Texte de cellule nettoyé ; chaîne vide pour les valeurs « XX »."""
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    texte = re.sub(r"\s+", " ", str(valeur)).strip()
    return "" if norm(texte) in VIDES else texte


def lire_date(valeur) -> date | None:
    """« 17 SEPTEMBRE 2025 », « 1ER MAI 2024 », 2025-09-17 ou une vraie date."""
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    texte = norm(valeur)
    m = re.search(r"(\d{1,2})(?:er)?\s+([a-z]+)\s+(\d{4})", texte)
    if m and m.group(2) in MOIS:
        return date(int(m.group(3)), MOIS[m.group(2)], int(m.group(1)))
    m = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", texte)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def separer_adresse(adresse: str) -> tuple[str, str]:
    """« 3600 BOULEVARD DU TRICENTENAIRE » → (« 3600 », « BOULEVARD DU TRICENTENAIRE »)."""
    m = re.match(r"\s*(\d+[A-Za-z]?)[\s,]+(.+)", adresse or "")
    if not m:
        return "", (adresse or "").strip()
    return m.group(1), m.group(2).strip()


def nom_personne(texte: str) -> str:
    """Retire la civilité : « M. ZACK DIOTTE » → « ZACK DIOTTE »."""
    return re.sub(r"^(m\.|mme|mlle|monsieur|madame)\s+", "", (texte or "").strip(), flags=re.I).strip()


def ouvrir_classeur(chemin):
    try:
        return load_workbook(chemin, data_only=True)
    except ValueError:
        # Certains classeurs déclarent une famille de police > 14, refusée par
        # openpyxl alors qu'Excel l'accepte : on relâche la borne et on relit.
        from openpyxl.styles.fonts import Font

        Font.__dict__["family"].max = 255
        return load_workbook(chemin, data_only=True)


class Feuille:
    """Accès par libellé à une feuille : grille {(ligne, colonne): valeur}."""

    def __init__(self, ws):
        self.cellules = {
            (c.row, c.column): c.value
            for ligne in ws.iter_rows()
            for c in ligne
            if c.value not in (None, "")
        }
        self.max_ligne = max((r for r, _c in self.cellules), default=0)

    def v(self, ligne, colonne):
        return self.cellules.get((ligne, colonne))

    def ligne(self, ligne) -> list[tuple[int, object]]:
        return sorted((c, v) for (r, c), v in self.cellules.items() if r == ligne)

    def trouver(self, libelle, depuis=1, jusqua=None, colonne=None) -> tuple[int, int] | None:
        """Première cellule dont le texte normalisé vaut `libelle` (ou commence
        par `libelle…` quand il finit par « … »)."""
        cible = norm(libelle.rstrip("…"))
        prefixe = libelle.endswith("…")
        for r in range(depuis, (jusqua or self.max_ligne) + 1):
            for c, v in self.ligne(r):
                if colonne is not None and c != colonne:
                    continue
                t = norm(v)
                if t == cible or (prefixe and t.startswith(cible)):
                    return r, c
        return None

    def valeur_apres(self, libelle, depuis=1, jusqua=None) -> str:
        """Valeur à droite du libellé, ou après le « : » dans la même cellule."""
        pos = self.trouver(libelle + "…", depuis, jusqua)
        if pos is None:
            return ""
        r, c = pos
        texte = str(self.v(r, c))
        if ":" in texte and texte.split(":", 1)[1].strip():
            return propre(texte.split(":", 1)[1])
        for c2, v in self.ligne(r):
            if c2 > c:
                return propre(v)
        return ""


def detecter_type(feuille: Feuille) -> str | None:
    texte = " ".join(norm(v) for (r, _c), v in feuille.cellules.items() if r <= 20)
    if "s-536" in texte or "s536" in texte:
        return ALARME
    if "nfpa 25" in texte:
        return GICLEURS
    if "eclairage" in texte and "nfpa 101" in texte:
        return ECLAIRAGE
    if "extincteur" in texte and "nfpa 10" in texte:
        return EXTINCTEURS
    return None


def lire_fichier(chemin) -> dict:
    """Point d'entrée : lit un classeur et renvoie son contenu selon son type."""
    wb = ouvrir_classeur(chemin)
    feuille = Feuille(wb.worksheets[0])
    type_ = detecter_type(feuille)
    lecteurs = {ALARME: lire_alarme, EXTINCTEURS: lire_extincteurs, ECLAIRAGE: lire_eclairage, GICLEURS: lire_gicleurs}
    if type_ is None:
        raise ValueError(f"Gabarit non reconnu : {chemin}")
    donnees = lecteurs[type_](feuille)
    donnees["type"] = type_
    donnees["fichier"] = str(chemin)
    return donnees


# ── Alarme S-536 ─────────────────────────────────────────────────────────────

COL_A, COL_B = 1, 2
# Cases à cocher : « X » en M = Oui, en O = Non, en Q = S.O.
CASES = {13: "oui", 15: "non", 17: "sans_objet"}
COLONNES_REPONSE = range(12, 18)  # L à Q
CODE_ITEM = re.compile(r"^[A-Z]{1,2}$")
UNITES = {"vcc", "vca", "v", "ma", "a", "ah"}


def _items(feuille: Feuille, debut: int, fin: int) -> list[tuple[str, int, int]]:
    """Items (code, première ligne, dernière ligne) d'une section du formulaire."""
    lignes = []
    for r in range(debut, fin):
        code = feuille.v(r, COL_A)
        if isinstance(code, str) and CODE_ITEM.match(code.strip()) and feuille.v(r, COL_B):
            lignes.append((code.strip(), r))
    return [(code, r, (lignes[i + 1][1] if i + 1 < len(lignes) else fin) - 1) for i, (code, r) in enumerate(lignes)]


def _reponse(feuille: Feuille, debut: int, fin: int) -> str | None:
    for r in range(debut, fin + 1):
        if norm(feuille.v(r, 12)) != "oui":
            continue
        for colonne, reponse in CASES.items():
            if norm(feuille.v(r, colonne)) == "x":
                return reponse
    return None


def _texte_mesure(feuille: Feuille, debut: int, fin: int) -> str:
    """Valeurs écrites dans les colonnes L à Q (tensions, dates…), libellés
    compris : « T : 26.55 VCC / C : … »."""
    morceaux = []
    for r in range(debut, fin + 1):
        cellules = [propre(feuille.v(r, c)) for c in COLONNES_REPONSE]
        if norm(cellules[0]) in ("oui", "non", "s.o"):
            continue
        valeurs = [t for t in cellules if t]
        # Une ligne sans valeur mesurée (« Heure : », « C : xx MA ») est ignorée.
        if any(not t.endswith(":") and norm(t) not in UNITES for t in valeurs):
            morceaux.append(" ".join(valeurs))
    return " / ".join(morceaux)


def _sections_alarme(feuille: Feuille) -> list[tuple[str, int, int]]:
    """Bornes des sections E1, E2.1… E2.12, E3 (repérées par leur titre en colonne A)."""
    titres = []
    for r in range(1, feuille.max_ligne + 1):
        t = norm(feuille.v(r, COL_A))
        m = re.match(r"^(e1|e2\.\d+|e3)\b", t)
        if m and (not titres or titres[-1][0] != m.group(1)):
            titres.append((m.group(1), r))
    return [(nom, r, (titres[i + 1][1] if i + 1 < len(titres) else feuille.max_ligne + 1)) for i, (nom, r) in enumerate(titres)]


def _texte_libre(feuille: Feuille, debut: int, fin: int, colonne_min=1) -> str:
    lignes = []
    for r in range(debut, fin):
        valeurs = [propre(v) for c, v in feuille.ligne(r) if c >= colonne_min]
        valeurs = [v for v in valeurs if v and not norm(v).startswith("(reference")]
        if valeurs:
            lignes.append(" ".join(valeurs))
    return "\n".join(lignes)


def lire_alarme(feuille: Feuille) -> dict:
    avertissements = []
    sections = {nom: (debut, fin) for nom, debut, fin in _sections_alarme(feuille)}

    # Bâtiment (fiche E1) : le nom, l'adresse et les deux lignes de ville dessous.
    adresse_pos = feuille.trouver("Adresse")
    adresse = ville = code_postal = ""
    if adresse_pos:
        r, _c = adresse_pos
        colonne = next((c for c, v in feuille.ligne(r) if c > adresse_pos[1]), None)
        adresse = propre(feuille.v(r, colonne)) if colonne else ""
        ville = propre(feuille.v(r + 1, colonne)).split(",")[0].strip() if colonne else ""
        code_postal = propre(feuille.v(r + 2, colonne)) if colonne else ""

    adresse_facturation = ", ".join(
        t for t in (feuille.valeur_apres("Adresse de facturation"), feuille.valeur_apres("Ville, Province, Code postal")) if t
    )
    donnees = {
        "client": {
            "nom": feuille.valeur_apres("Nom du commerce ou propriétaire"),
            "contact_nom": feuille.valeur_apres("À l'attention de"),
            "contact_telephone": feuille.valeur_apres("Téléphone"),
            "contact_email": feuille.valeur_apres("Courriel"),
            "adresse": adresse_facturation,
        },
        "batiment": {
            "nom": feuille.valeur_apres("Nom du bâtiment"),
            "adresse": adresse,
            "ville": ville,
            "code_postal": code_postal,
            "fabricant_reseau": feuille.valeur_apres("Fabricant du réseau"),
            "modele_systeme": feuille.valeur_apres("Modèle"),
        },
        "date": lire_date(feuille.valeur_apres("Date")),
        "technicien": "",
        "e1": {},
        "e2": {},
        "legende": {},
        "dispositifs": [],
        "avertissements": avertissements,
    }

    # E1 — réponses A à H, commentaires (G) et technicien signataire.
    if "e1" in sections:
        debut, fin = sections["e1"]
        for code, r1, r2 in _items(feuille, debut, fin):
            if code == "G":
                commentaire = _texte_libre(feuille, r1, r2 + 1, colonne_min=3)
                donnees["e1"]["G"] = commentaire.split(":", 1)[-1].strip() if norm(commentaire).startswith("commentaires") else commentaire
            elif code <= "H":
                donnees["e1"][code] = _reponse(feuille, r1, r2)
        pos = feuille.trouver("Nom en lettres moulées du technicien…", debut, fin)
        if pos:
            donnees["technicien"] = propre(feuille.v(pos[0] - 1, COL_A))

    # E2.x — même format que FicheE2.details : {"e2_1": {"localisation", "description", "A": "oui", …}}.
    for nom, (debut, fin) in sections.items():
        if not nom.startswith("e2."):
            continue
        cle = nom.replace(".", "_")
        section = {}
        emplacement = feuille.trouver("Emplacement…", debut, fin, colonne=COL_A)
        if emplacement:
            section["localisation"] = feuille.valeur_apres("Emplacement", emplacement[0], emplacement[0])
        identification = feuille.trouver("Identification…", debut, fin, colonne=COL_A)
        if identification:
            section["description"] = feuille.valeur_apres("Identification", identification[0], identification[0])
        if cle == "e2_12":
            remarques = _texte_libre(feuille, debut + 1, fin)
            if remarques:
                section["remarques"] = remarques
        for code, r1, r2 in _items(feuille, debut, fin):
            reponse = _reponse(feuille, r1, r2)
            section[code] = reponse if reponse else _texte_mesure(feuille, r1, r2)
        donnees["e2"][cle] = {k: v for k, v in section.items() if v}

    # Légende (TYPE / NO DE MODÈLE / QUANTITÉ) et fiche des dispositifs E3.2.1.
    debut_e3, fin_e3 = sections.get("e3", (0, 0))
    entete = None
    for r in range(debut_e3, fin_e3):
        if norm(feuille.v(r, COL_A)) == "type" and any(norm(v) == "quantite" for _c, v in feuille.ligne(r)):
            entete = r
            break
    if entete:
        col_qte = next(c for c, v in feuille.ligne(entete) if norm(v) == "quantite")
        r = entete + 1
        while r < fin_e3 and propre(feuille.v(r, COL_A)) and not norm(feuille.v(r, COL_A)).startswith("notes"):
            donnees["legende"][propre(feuille.v(r, COL_A)).upper()] = {
                "description": propre(feuille.v(r, COL_B)),
                "quantite": propre(feuille.v(r, col_qte)),
            }
            r += 1

    colonnes = None
    for r in range(debut_e3, fin_e3):
        entetes = {norm(v): c for c, v in feuille.ligne(r)}
        if "emplacement" in entetes and "dispositif" in entetes:
            colonnes = {
                "emplacement": entetes["emplacement"],
                "dispositif": entetes["dispositif"],
                "installation": next(c for t, c in entetes.items() if t.startswith("installation")),
                "entretien": next(c for t, c in entetes.items() if t.startswith("necessite")),
                "alarme": next(c for t, c in entetes.items() if t.startswith("fonctionnement")),
                "annonce": next(c for t, c in entetes.items() if t.startswith("indication")),
                "zone": next((c for t, c in entetes.items() if t.startswith("zone")), None),
                "remarque": next((c for t, c in entetes.items() if t.startswith("remarque")), None),
            }
            continue
        if colonnes is None:
            continue
        emplacement = propre(feuille.v(r, colonnes["emplacement"]))
        code = propre(feuille.v(r, colonnes["dispositif"])).upper()
        if not emplacement or not code or emplacement.startswith("…"):
            continue
        donnees["dispositifs"].append({
            "ligne": r,
            "emplacement": emplacement,
            "code": code,
            "installation": norm(feuille.v(r, colonnes["installation"])),
            "entretien": norm(feuille.v(r, colonnes["entretien"])),
            "alarme": norm(feuille.v(r, colonnes["alarme"])),
            "annonce": norm(feuille.v(r, colonnes["annonce"])),
            "zone": propre(feuille.v(r, colonnes["zone"])) if colonnes["zone"] else "",
            "remarque": propre(feuille.v(r, colonnes["remarque"])) if colonnes["remarque"] else "",
        })

    if donnees["date"] is None:
        avertissements.append("Date d'inspection illisible.")
    return donnees


# ── Extincteurs NFPA 10 ──────────────────────────────────────────────────────

COLONNES_EXTINCTEURS = {
    "numero": "numero",
    "derniere inspection": "derniere_inspection",
    "numero de serie": "numero_serie",
    "poids": "poids",
    "type": "type",
    "classe": "classe",
    "annee de fabrication": "annee_fabrication",
    "derniere recharge": "derniere_recharge",
    "dernier test hydro": "dernier_test_hydro",
    "prochaine recharge": "prochaine_recharge",
    "prochain test hydro": "prochain_test_hydro",
    "etiquette de visibilite": "etiquette",
    "support murale": "support",
    "lieu extincteur": "lieu",
}


def lire_extincteurs(feuille: Feuille) -> dict:
    donnees = {
        "client": {
            "nom": feuille.valeur_apres("Nom"),
            "contact_nom": feuille.valeur_apres("Responsable"),
            "contact_telephone": feuille.valeur_apres("# Téléphone"),
        },
        "batiment": {"nom": feuille.valeur_apres("Nom"), "adresse": feuille.valeur_apres("Adresse de service")},
        "technicien": feuille.valeur_apres("Inspection faite par"),
        "date": lire_date(feuille.valeur_apres("Date")),
        "extincteurs": [],
        "avertissements": [],
    }
    colonnes = None
    for r in range(1, feuille.max_ligne + 1):
        entetes = {norm(v): c for c, v in feuille.ligne(r)}
        if "numero" in entetes and "lieu extincteur" in entetes:
            colonnes = {COLONNES_EXTINCTEURS[t]: c for t, c in entetes.items() if t in COLONNES_EXTINCTEURS}
            continue
        if norm(feuille.v(r, COL_A)).startswith("remarques"):
            colonnes = None  # fin des tableaux
        if colonnes is None:
            continue
        ligne = {cle: propre(feuille.v(r, c)) for cle, c in colonnes.items()}
        if not any(v for cle, v in ligne.items() if cle != "numero"):
            continue  # ligne numérotée mais vide du gabarit
        ligne["ligne"] = r
        donnees["extincteurs"].append(ligne)
    if donnees["date"] is None:
        donnees["avertissements"].append("Date d'inspection illisible.")
    donnees["notes"] = _notes_bas_de_page(feuille)
    return donnees


# ── Éclairage d'urgence NFPA 101 ─────────────────────────────────────────────

def lire_eclairage(feuille: Feuille) -> dict:
    nombre = feuille.valeur_apres("Nombre de dispositif")
    donnees = {
        "client": {
            "nom": feuille.valeur_apres("Nom du bâtiment"),
            "contact_nom": feuille.valeur_apres("Nom du contact"),
            "contact_telephone": feuille.valeur_apres("Téléphone du contact"),
            "contact_email": feuille.valeur_apres("Courriel"),
        },
        "batiment": {"nom": feuille.valeur_apres("Nom du bâtiment"), "adresse": feuille.valeur_apres("Adresse de service")},
        "technicien": feuille.valeur_apres("Nom du technicien"),
        "date": lire_date(feuille.valeur_apres("Date de l'inspection")),
        "nombre_annonce": int(nombre) if nombre.isdigit() else None,
        "unites": [],
        "avertissements": [],
    }
    colonnes = None
    for r in range(1, feuille.max_ligne + 1):
        entetes = {norm(v): c for c, v in feuille.ligne(r)}
        if "no" in entetes and "emplacement" in entetes:
            colonnes = {
                "emplacement": entetes["emplacement"],
                "fabricant": entetes.get("fabricant"),
                "dispositif": entetes.get("dispositif"),
                "modele_batterie": entetes.get("modele"),
                # Selon la copie du gabarit : « FONCTIONNEL » ou « CONFORME ».
                "fonctionnel": entetes.get("fonctionnel") or entetes.get("conforme"),
                "capacite": entetes.get("capacite"),
            }
            continue
        if norm(feuille.v(r, COL_A)).startswith("legende"):
            break
        if colonnes is None:
            continue
        ligne = {cle: propre(feuille.v(r, c)) if c else "" for cle, c in colonnes.items()}
        if not ligne["emplacement"]:
            continue
        ligne["ligne"] = r
        donnees["unites"].append(ligne)
    if donnees["date"] is None:
        donnees["avertissements"].append("Date d'inspection illisible.")
    if donnees["nombre_annonce"] is not None and donnees["nombre_annonce"] != len(donnees["unites"]):
        donnees["avertissements"].append(
            f"Le rapport annonce {donnees['nombre_annonce']} dispositif(s) mais le tableau en contient {len(donnees['unites'])}."
        )
    donnees["notes"] = _notes_bas_de_page(feuille)
    return donnees


# ── Gicleurs NFPA 25 ─────────────────────────────────────────────────────────

REPONSES_GICLEUR = {"oui": "oui", "non": "non", "n/a": "na"}
SYSTEMES_GICLEUR = ("sous eau", "sous air", "preaction", "preaction electrique", "pompe")
VANNES_GICLEUR = ("papillon", "os&y", "piv", "murale")


def _est_coche(valeur) -> bool:
    return norm(valeur) == "x"


def _voisin(feuille: Feuille, r: int, c: int) -> str:
    """Valeur à droite d'un libellé — vide si la cellule suivante est un
    autre libellé (« OS&Y : », « Pression d'arrêt (lb/po²) »)."""
    for c2, v in feuille.ligne(r):
        if c2 > c:
            if isinstance(v, str) and (":" in v or "lb/po" in norm(v)):
                return ""
            return propre(v)
    return ""


def _apres(feuille: Feuille, libelle: str, depuis=1, jusqua=None) -> str:
    pos = feuille.trouver(libelle + "…", depuis, jusqua)
    return _voisin(feuille, *pos) if pos else ""


def _coches_en_ligne(feuille: Feuille, r: int) -> set[str]:
    """Libellés suivis de « : X » dans le texte de la ligne (« Oui : X »,
    « Scellés : X   Barrées : X »)."""
    texte = norm(" | ".join(str(v) for _c, v in feuille.ligne(r)))
    return {m.strip() for m in re.findall(r"([a-z/&.\- ]+?)\s*:\s*x(?![a-z])", texte)}


def _choix_coche(feuille: Feuille, r: int) -> str:
    """Libellé le plus proche à gauche du « X » de la ligne (« Annuelle  X »)."""
    libelle = ""
    for _c, v in feuille.ligne(r):
        if _est_coche(v):
            return libelle
        if isinstance(v, str) and norm(v):
            libelle = propre(v).rstrip(" :")
    return ""


def _duree(valeur) -> dict | None:
    """Délai « minutes/secondes ». Excel range « 0:30 » en h:mm : on le lit
    comme 0 min 30 s, seule lecture plausible pour ce délai."""
    if isinstance(valeur, datetime):
        valeur = valeur.time()
    if isinstance(valeur, time):
        if valeur.second:
            return {"min": str(valeur.hour * 60 + valeur.minute), "sec": str(valeur.second)}
        return {"min": str(valeur.hour), "sec": str(valeur.minute)}
    texte = propre(valeur)
    m = re.fullmatch(r"(\d+)\s*[:m]\s*(\d+)\s*s?", norm(texte))
    if m:
        return {"min": m.group(1), "sec": m.group(2)}
    return {"min": texte, "sec": ""} if texte else None


def _notes_bas_de_page(feuille: Feuille) -> dict:
    """Remarques / Recommandations / Obligations du bas des gabarits
    extincteurs et éclairage."""
    return {
        "remarques": _texte_entre(feuille, "Remarques", "Recommandations"),
        "recommandations": _texte_entre(feuille, "Recommandations", "Obligations"),
        "obligations": _texte_entre(feuille, "Obligations", None),
    }


def _texte_entre(feuille: Feuille, titre: str, fin: str | None) -> str:
    """Texte libre sous un titre (« REMARQUES : »), jusqu'au titre suivant."""
    debut = feuille.trouver(titre, colonne=COL_A)
    if debut is None:
        return ""
    arret = feuille.trouver(fin + "…", debut[0] + 1, colonne=COL_A) if fin else None
    suite = " ".join(propre(v) for c, v in feuille.ligne(debut[0]) if c > debut[1] and propre(v))
    texte = _texte_libre(feuille, debut[0] + 1, arret[0] if arret else feuille.max_ligne + 1)
    return "\n".join(t for t in (suite, texte) if t)


def _pompe_gicleur(feuille: Feuille, titre: str, avertissements: list) -> dict:
    """Bloc « Pompe de surpression » / « Pompe incendie » : présence et mode
    cochés dans le texte (« Oui : X »), pressions et délai."""
    pos = feuille.trouver(titre, colonne=COL_A)
    if pos is None:
        return {}
    r0 = pos[0]
    ligne_oui = feuille.trouver("Oui…", r0, r0 + 4, colonne=COL_A)
    coches = _coches_en_ligne(feuille, ligne_oui[0]) if ligne_oui else set()
    pompe = {
        "presente": "oui" if "oui" in coches else "non" if "non" in coches else "",
        "mode": next((m for m in ("manuelle", "automatique") if m in coches), ""),
        "depart": _apres(feuille, "Pression de départ", r0, r0 + 5),
        "arret": _apres(feuille, "Pression d'arrêt", r0, r0 + 5),
        "delai": None,
    }
    ligne_delai = feuille.trouver("Délai…", r0, r0 + 6)
    if ligne_delai:
        brut = next((v for c, v in feuille.ligne(ligne_delai[0]) if c > ligne_delai[1]), None)
        pompe["delai"] = _duree(brut)
        if isinstance(brut, time) and not brut.second:
            avertissements.append(
                f"{titre} : délai saisi « {brut:%H:%M} » (heures:minutes dans Excel), "
                f"lu comme {pompe['delai']['min']} min {pompe['delai']['sec']} s."
            )
    return pompe


def lire_gicleurs(feuille: Feuille) -> dict:
    pos_date = feuille.trouver("Date de l'inspection…")
    donnees = {
        "client": {
            "nom": feuille.valeur_apres("Nom du bâtiment"),
            "contact_nom": feuille.valeur_apres("Nom du contact"),
            "contact_telephone": feuille.valeur_apres("Numéro de téléphone"),
        },
        "batiment": {"nom": feuille.valeur_apres("Nom du bâtiment"), "adresse": feuille.valeur_apres("Adresse")},
        "technicien": feuille.valeur_apres("Nom du technicien"),
        "date": None,
        "frequence": "",
        "synthese": {},
        "systemes": {},
        "vannes": {},
        "questions": [],
        "valeurs": {},
        "supervisions": [],
        "avertissements": [],
    }
    if pos_date:
        # Vraie date Excel à droite du libellé, ou texte « 14 OCTOBRE 2025 ».
        brut = next((v for c, v in feuille.ligne(pos_date[0]) if c > pos_date[1]), None)
        donnees["date"] = lire_date(brut) or lire_date(feuille.valeur_apres("Date de l'inspection"))

    pos = feuille.trouver("Type d'inspection…")
    if pos:
        donnees["frequence"] = _choix_coche(feuille, pos[0])

    for cle, libelle in (
        ("fonctionnel", "Suite à l'inspection, le système est fonctionnel"),
        ("deficiences", "Le système présente des déficiences"),
        ("recommandations", "Nous ajoutons des recommandations"),
        ("relie_alarme", "Relié à l'alarme incendie"),
    ):
        donnees["synthese"][cle] = REPONSES_GICLEUR.get(norm(_apres(feuille, libelle)), "")

    for systeme in SYSTEMES_GICLEUR:
        pos = feuille.trouver(systeme, colonne=COL_A)
        if pos and _voisin(feuille, *pos):
            donnees["systemes"][systeme] = _voisin(feuille, *pos)
    pos = feuille.trouver("Papillon…")
    if pos:
        for c, v in feuille.ligne(pos[0]):
            if norm(v) in VANNES_GICLEUR and _voisin(feuille, pos[0], c):
                donnees["vannes"][norm(v)] = _voisin(feuille, pos[0], c)

    # Questions Oui / Non / N/A : colonnes données par la ligne d'en-tête
    # qui précède chaque bloc. Une ligne qui porte d'autres libellés avant
    # les cases (« Interrupteur d'alarme … Interrupteur de débit X ») est un
    # choix, pas une question.
    colonnes = {}
    for r in range(1, feuille.max_ligne + 1):
        cellules = feuille.ligne(r)
        entetes = {c: REPONSES_GICLEUR[norm(v)] for c, v in cellules if norm(v) in REPONSES_GICLEUR}
        if len(entetes) >= 2:
            colonnes = entetes
            continue
        coches = {c for c, v in cellules if _est_coche(v)}
        libelle = propre(feuille.v(r, COL_A))
        if not coches or not libelle or not colonnes or not coches <= set(colonnes):
            continue
        if any(COL_A < c < min(colonnes) and isinstance(v, str) and norm(v) for c, v in cellules):
            continue
        reponses = {colonnes[c] for c in coches}
        if len(reponses) > 1:
            donnees["avertissements"].append(f"Ligne {r} « {libelle[:60]} » : plusieurs cases cochées, ignorée.")
            continue
        donnees["questions"].append({"libelle": libelle, "reponse": reponses.pop(), "ligne": r})

    v = donnees["valeurs"]
    pos = feuille.trouver("Scellés…")
    v["etat_vannes"] = sorted(_coches_en_ligne(feuille, pos[0])) if pos else []
    v["nombre_raccords"] = _apres(feuille, "Nombre de raccords")
    debut = (feuille.trouver("Système de gicleurs sous eau…") or (1, 1))[0]
    for cle, libelle in (
        ("numero_systeme", "# du système"), ("marque", "Marque"), ("modele", "Modèle"), ("annee", "Année"),
        ("no_serie", "No de série"), ("diametre", "Diamètre"), ("salle", "Emplacement de la salle des gicleurs"),
        ("zone", "Zone protégée"), ("robinet_essai", "Localisation du robinet d'essai"),
        ("pression_systeme", "Pression du système avant les essais"),
    ):
        v[cle] = _apres(feuille, libelle, depuis=debut)
    pos = feuille.trouver("Interrupteur d'alarme…")
    v["interrupteur"] = _choix_coche(feuille, pos[0]) if pos else ""

    donnees["pompe_surpression"] = _pompe_gicleur(feuille, "Pompe de surpression", donnees["avertissements"])
    donnees["pompe_incendie"] = _pompe_gicleur(feuille, "Pompe incendie", donnees["avertissements"])

    debut = (feuille.trouver("Système antigel…") or (1, 1))[0]
    for cle, libelle in (("antigel_secteur", "Secteur protégé"), ("antigel_localisation", "Localisation du système antigel"),
                         ("antigel_type", "Type"), ("antigel_resultat", "Résultat"), ("antigel_temperature", "Température")):
        v[cle] = _apres(feuille, libelle, debut, debut + 8)

    debut = (feuille.trouver("Vérification de l'alimentation d'eau…") or (1, 1))[0]
    for cle, libelle in (("alim_ville", "Ville"), ("alim_diametre", "Diamètre"), ("alim_statique", "Pression statique"),
                         ("alim_pompe", "Pompe incendie"), ("alim_reservoir", "Réservoir"),
                         ("alim_residuelle", "Pression résiduelle"), ("alim_ecart_10", "Est-ce que le résultat indique")):
        v[cle] = _apres(feuille, libelle, debut, debut + 12)

    pos = feuille.trouver("Liste des supervisions…")
    entete = feuille.trouver("Emplacement", pos[0], pos[0] + 3, colonne=COL_A) if pos else None
    if entete:
        cols = {norm(val): c for c, val in feuille.ligne(entete[0])}
        fin = feuille.trouver("Remarques…", entete[0] + 1, colonne=COL_A)
        for r in range(entete[0] + 1, fin[0] if fin else feuille.max_ligne + 1):
            ligne = {
                "emplacement": propre(feuille.v(r, COL_A)),
                "type": propre(feuille.v(r, cols.get("type de valves", 0))),
                "diametre": propre(feuille.v(r, cols.get("diametre", 0))),
                "fonctionnement": propre(feuille.v(r, cols.get("bon fonctionnement", 0))),
            }
            if any(ligne.values()):
                donnees["supervisions"].append(ligne)

    donnees["textes"] = {
        "observations": _texte_entre(feuille, "Remarques et observations", "Type de système et nombre"),
        "remarques": _texte_entre(feuille, "Remarques", "Recommandations"),
        "recommandations": _texte_entre(feuille, "Recommandations", "Obligations"),
        "obligations": _texte_entre(feuille, "Obligations", None),
    }

    if donnees["date"] is None:
        donnees["avertissements"].append("Date d'inspection illisible.")
    if not donnees["questions"]:
        donnees["avertissements"].append("Aucune case Oui / Non / N/A lue.")
    return donnees
